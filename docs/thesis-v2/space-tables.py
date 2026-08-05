#!/usr/bin/env python3
"""Nới giãn cách cho bảng trong báo cáo thesis-v2 (chạy TRƯỚC inject-cover).

Pandoc render ô bảng với style `Compact` (giãn dòng đơn 1.0) và style bảng có
padding trên/dưới = 0 → chữ trong ô dán sát mép và các dòng sát nhau. Script này
hậu xử lý docx do pandoc sinh ra:
  1. Thêm padding trên/dưới cho mọi ô bảng (breathing room theo chiều dọc).
  2. Nới giãn dòng + khoảng cách đoạn cho chữ bên trong ô.
Chỉ áp cho các bảng nội dung; trang bìa được inject sau nên không bị ảnh hưởng.
"""
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CELL_TOP = 55       # twips (~0.1cm) — đủ thoáng, không phồng trang
CELL_BOTTOM = 55
CELL_LEFT = 115
CELL_RIGHT = 115
LINE_SPACING = 1.15  # giãn dòng trong ô (body ngoài bảng là 1.5)


def set_table_cell_margins(table):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblCellMar'))
    if old is not None:
        tblPr.remove(old)
    mar = OxmlElement('w:tblCellMar')
    for tag, val in [('w:top', CELL_TOP), ('w:left', CELL_LEFT),
                     ('w:bottom', CELL_BOTTOM), ('w:right', CELL_RIGHT)]:
        e = OxmlElement(tag)
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblPr.append(mar)


def loosen_cell_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.line_spacing = LINE_SPACING
                pf.space_before = Pt(2)
                pf.space_after = Pt(2)


def enable_hyphenation(doc):
    """Bật tự động ngắt từ để giảm khoảng trắng do căn đều 2 bên.

    Ngôn ngữ mặc định là en-US nên chỉ các từ tiếng Anh (Spring, Application…)
    bị ngắt; từ tiếng Việt không có trong từ điển Anh nên không bị ngắt bậy.
    doNotHyphenateCaps tránh ngắt từ viết hoa (acronym).
    """
    settings = doc.settings.element
    for tag, val in [('w:autoHyphenation', 'true'),
                     ('w:doNotHyphenateCaps', 'true'),
                     ('w:hyphenationZone', '357')]:  # ~0.63cm
        if settings.find(qn(tag)) is None:
            el = OxmlElement(tag)
            el.set(qn('w:val'), val)
            settings.append(el)


def main(path):
    doc = Document(path)
    n = 0
    for table in doc.tables:
        set_table_cell_margins(table)
        loosen_cell_paragraphs(table)
        n += 1
    enable_hyphenation(doc)
    doc.save(path)
    print(f"space-tables: nới giãn cách {n} bảng + bật hyphenation")


if __name__ == '__main__':
    main(sys.argv[1])
